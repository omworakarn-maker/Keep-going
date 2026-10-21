from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path
from copy import deepcopy
f=next(Path('.').glob('*Mobile*.docx'));d=Document(f)
body=d._element.body
start=next(e for e in body if ''.join(e.itertext()).startswith('4.1 ภาพ'))
end=next(e for e in body if ''.join(e.itertext()).startswith('5. อภิปราย'))
e=start
while e is not end:
 nxt=e.getnext();body.remove(e);e=nxt
items=[('01-journal','หน้าบันทึกประจำวัน','ผู้ใช้เลือกเช็กอินหรือไดอารี เปิดพื้นที่เขียนจดหมาย และอ่านบันทึกของวันนี้ได้ สถานะในภาพแสดงว่าบัญชีนี้เช็กอินแล้วในวันดังกล่าว'),('02-history','หน้าประวัติย้อนหลัง','รายการที่บันทึกแสดงประเภท อารมณ์ วันที่ และข้อความ ผู้ใช้เลือกตัวกรองเพื่อดูเฉพาะเช็กอิน ไดอารี หรือจดหมายที่เก็บไว้ได้'),('03-progress','หน้าการเติบโต','หน้าสรุปแสดงจำนวนวัน No Contact จำนวนกิจกรรม และเส้นทางอารมณ์ตามช่วงเวลา ตัวเลขในภาพเป็นสถานะของบัญชีขณะบันทึกภาพ'),('04-settings','หน้าตั้งค่า','ผู้ใช้ตรวจสอบบัญชี กำหนดวันเริ่ม No Contact และเพิ่มเหตุผลที่เลือกตัวเอง ข้อมูลเหล่านี้เป็นส่วนหนึ่งของสถานะผู้ใช้ในแอป'),('05-urge-support','หน้าเครื่องมือพักใจ','กระบวนการพักใจแสดงข้อความแนะนำและเวลานับถอยหลังสำหรับฝึกหายใจ พร้อมปุ่มไปเขียนความรู้สึก โดยผู้ใช้ย้อนกลับหรือปิดเครื่องมือได้'),('06-release-letter','หน้าจดหมายที่ไม่ต้องส่ง','ผู้ใช้เขียนข้อความ ระบุผู้รับได้ตามต้องการ และเลือกเก็บไว้ในประวัติหรือปล่อยจดหมายไป ทั้งสองทางเลือกไม่ส่งข้อความถึงบุคคลอื่น')]
def para(text='',bold=False,size=14):
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(5)
 r=p.add_run(text);r.font.name='TH SarabunPSK';r.font.size=Pt(size);r.bold=bold
 end.addprevious(p._p);return p
for i,(name,title,desc) in enumerate(items,1):
 p=para(('4.1 ภาพหน้าจอการใช้งานแอป' if i==1 else '4.1 ภาพหน้าจอการใช้งานแอป ต่อ'),True,16)
 p.paragraph_format.page_break_before=True;p.paragraph_format.keep_with_next=True
 if i==1:
  p=para('ภาพที่ 1–6 บันทึกจากแอป Keep Going ที่ทำงานผ่าน Expo Go บน iPhone Simulator โดยวงเล็บและหมายเลขในภาพอธิบายส่วนติดต่อที่สำคัญ');p.paragraph_format.keep_with_next=True
 p=para();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
 r=p.add_run();r.add_picture('docs/screenshots/annotated/'+name+'-annotated.png',width=Inches(5.25))
 for shape in p._p.xpath('.//wp:docPr'):shape.set('descr',title+' พร้อมวงเล็บและคำอธิบายภาษาไทย')
 p=para('ภาพที่ '+str(i)+' '+title,True);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
 para(desc)
# Keep discussion separate from last figure.
from docx.text.paragraph import Paragraph
Paragraph(end,d._body).paragraph_format.page_break_before=True
d.save(f)
Path('docs/word-qa/mobile.docx').write_bytes(f.read_bytes())
print(f)
