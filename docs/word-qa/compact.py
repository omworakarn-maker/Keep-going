from docx import Document
from pathlib import Path
f=next(Path('.').glob('*Mobile*.docx'));d=Document(f)
a=next(i for i,p in enumerate(d.paragraphs) if p.text.startswith('3.1 '));b=next(i for i,p in enumerate(d.paragraphs) if p.text.startswith('4. ผล'))
texts=[
'โครงงานประยุกต์ React Native Paper ใน MySpaceTools ของหน้าตั้งค่า โดยใช้ PaperProvider กำหนดธีมร่วม อ่านข้อมูลจาก AppContext และใช้ Expo Router เปิดหน้าปลายทาง IconButton และ Tooltip ใช้เปิดเมนูพร้อมคำอธิบาย Menu และ Drawer เป็นทางลัดไปยังหน้าหลัก ส่วน List และ Divider จัดกลุ่มข้อมูล',
'DataTable แสดงจำนวนเช็กอิน ไดอารี และข้อความที่ไม่ได้ส่ง โดยแตะแถวเพื่อเปิดประวัติ Button เปิด Dialog ผ่าน Portal เพื่ออ่าน selfReasons รองรับกรณีว่างและข้อความยาว ส่วน Switch ควบคุมสรุปก่อนวันอาทิตย์และการแจ้งเตือนผ่านบริการเดิม ตัวเลขอาจมาจากสำเนาในอุปกรณ์ระหว่างซิงก์ จึงไม่ใช่หลักฐานยืนยันข้อมูลเซิร์ฟเวอร์แบบเรียลไทม์',
'การตรวจรับครอบคลุมจำนวนรายการหลังเพิ่มบันทึก ทางลัด การเปิดปิด Dialog การเลื่อนอ่าน และการคงค่า Switch หลังเปิดแอปใหม่ โดยยังต้องทดสอบปฏิสัมพันธ์บนโทรศัพท์จริงเพิ่มเติม']
ps=d.paragraphs[a+1:b]
for p,t in zip(ps,texts):
 r=p.runs[0];r.text=t
 for extra in p.runs[1:]:extra._element.getparent().remove(extra._element)
for p in ps[len(texts):]:p._p.getparent().remove(p._p)
d.save(f);Path('docs/word-qa/mobile.docx').write_bytes(f.read_bytes())
