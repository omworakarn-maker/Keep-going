from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
f=next(Path('.').glob('*Cloud*.docx'));d=Document(f)
anchor=next(p for p in d.paragraphs if p.text.startswith('5. อภิปราย'))
items=[
('01-client','การเชื่อมต่อ Supabase','src/services/supabase.ts บรรทัด 1–19',[
'บรรทัด 5–6 อ่าน Project URL และ publishable key จากตัวแปรสภาพแวดล้อม บรรทัด 8–10 ตรวจว่ากำหนดค่าครบก่อนสร้าง client ภาพแสดงเฉพาะชื่อตัวแปรโดยไม่มีค่ากุญแจหรือรหัสผ่าน',
'บรรทัด 12 สร้าง supabase client ที่บริการข้อมูลนำไปใช้ร่วมกัน บรรทัด 14–17 กำหนดให้ AsyncStorage เก็บ session ต่ออายุโทเคนอัตโนมัติ และคงการเข้าสู่ระบบเมื่อเปิดแอปอีกครั้ง ส่วนนี้เป็นการตั้งค่าการเชื่อมต่อและ Auth ยังไม่ใช่คำสั่งเพิ่มแถวในฐานข้อมูล']),
('02-schema','โครงสร้างตารางและสิทธิ์ข้อมูล','supabase/migrations/20261021062416_create_daily_checkins.sql บรรทัด 1–30',[
'บรรทัด 1–11 สร้างตาราง daily_checkins โดย id เป็น UUID และ user_id อ้างอิงบัญชีใน auth.users ฟิลด์ checkin_date, mood, emoji และ note เก็บวันที่ อารมณ์ อีโมจิ และข้อความ เงื่อนไข unique (user_id, checkin_date) จำกัดให้แต่ละบัญชีมีหนึ่งแถวต่อวัน',
'บรรทัด 16 เปิด Row Level Security และบรรทัด 17 ให้สิทธิ์ CRUD แก่ผู้ที่เข้าสู่ระบบ นโยบายในบรรทัด 19–30 ตรวจ auth.uid() เทียบ user_id เพื่อจำกัดการอ่าน เพิ่ม และแก้ไขเฉพาะข้อมูลของตนเอง ทั้ง using และ with check ใช้ตรวจแถวตามการดำเนินการ ส่วน policy การลบอยู่ถัดจากช่วงโค้ดในภาพ']),
('03-checkin','คำสั่งบันทึกเช็กอินลงฐานข้อมูล','src/services/dailyCheckins.ts บรรทัด 25–43',[
'บรรทัด 26–27 ตรวจผู้ใช้ปัจจุบันด้วย supabase.auth.getUser() หากไม่พบบัญชีหรือเกิดข้อผิดพลาดจะหยุดการทำงาน จากนั้นเลือกตาราง daily_checkins และส่ง user_id พร้อมวันที่ อารมณ์ อีโมจิ ข้อความ และเวลาแก้ไข',
'คำสั่ง upsert ใช้ onConflict: user_id,checkin_date ให้ตรงกับ unique constraint ใน SQL ถ้ามีแถววันเดียวกันของบัญชีนั้นจะอัปเดตแถวเดิม หากยังไม่มีจะเพิ่มใหม่ select และ single ขอข้อมูลแถวที่บันทึกกลับมา ส่วน rowToEntry แปลงชื่อฟิลด์ฐานข้อมูลเป็นรูปแบบที่แอปใช้งาน ฟังก์ชันโยน error ให้ผู้เรียกจัดการเมื่อบันทึกไม่สำเร็จ']),
('04-diary','การเพิ่มและลบไดอารี','src/services/cloudData.ts บรรทัด 100–118',[
'ฟังก์ชัน upsertDiaryEntry อ่านรหัสผู้ใช้แล้วส่งข้อมูลไปตาราง diary_entries โดยจับคู่ id, createdAt, title, content, mood และ emoji กับคอลัมน์ฐานข้อมูล กำหนดอารมณ์และอีโมจิที่ไม่มีค่าเป็น null ตาม schema',
'onConflict: user_id,id ใช้คู่รหัสบัญชีและรหัสรายการเพื่อเพิ่มหรือแก้ไขรายการเดิม ส่วน deleteCloudDiaryEntry ใช้ delete ร่วมกับ eq สองเงื่อนไข ให้ลบเฉพาะ id ที่ระบุและเป็นของบัญชีปัจจุบัน การกรองจาก client ทำงานร่วมกับ RLS ที่ฐานข้อมูล มิใช่ใช้แทนนโยบายสิทธิ์']),
('05-context','จุดเชื่อมจากแอปไปยังบริการฐานข้อมูล','src/context/AppContext.tsx บรรทัด 109–132',[
'addEntry สร้างข้อมูลเช็กอินและปรับ state ในแอปก่อน เมื่อมีผู้ใช้ที่เข้าสู่ระบบจึงเรียก upsertDailyCheckin จากภาพที่ 3 หลังเซิร์ฟเวอร์ตอบกลับ ระบบนำ savedEntry มาแทนข้อมูลวันเดียวกัน ทำให้ใช้ id และข้อมูลที่ฐานข้อมูลส่งกลับ',
'addDiaryEntry เตรียมหัวข้อ เนื้อหา และเวลา เพิ่มรายการเข้า state แล้วเรียก upsertDiaryEntry จากภาพที่ 4 เส้นทางการบันทึกจึงเป็นหน้าจอเรียก Context จากนั้น Context เรียก service และ service ส่งข้อมูลไป Supabase',
'ข้อจำกัดคือ catch(() => undefined) ในภาพละทิ้งข้อผิดพลาดของการส่งข้อมูล ผู้ใช้จึงอาจเห็นรายการใน state แม้การบันทึกบนคลาวด์ล้มเหลว ภาพนี้ยืนยันจุดเชื่อมในโค้ด แต่ไม่ยืนยันว่าทุกรายการส่งสำเร็จ ควรเพิ่มสถานะ pending, synced และ failed พร้อมการแจ้งเตือนหรือส่งซ้ำ'])]
def p(text='',bold=False,size=14):
 q=d.add_paragraph();q.paragraph_format.space_after=Pt(6)
 r=q.add_run(text);r.bold=bold;r.font.name='TH SarabunPSK';r.font.size=Pt(size)
 anchor._p.addprevious(q._p);return q
for i,(name,title,source,notes) in enumerate(items,1):
 q=p('4.1 โค้ดที่ใช้เชื่อมต่อและจัดการฐานข้อมูล'+(' ต่อ' if i>1 else ''),True,16);q.paragraph_format.page_break_before=True;q.paragraph_format.keep_with_next=True
 q=p();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.keep_with_next=True
 q.add_run().add_picture('docs/code-figures/'+name+'.png',width=Inches(6.15))
 q=p('ภาพที่ '+str(i)+' '+title,True);q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.keep_with_next=True
 q=p('ที่มา '+source,size=12)
 for note in notes:p(note)
anchor.paragraph_format.page_break_before=True
d.save(f);Path('docs/word-qa/cloud.docx').write_bytes(f.read_bytes());print(f)
